"""Insert print figures with their captions kept on the same page."""
from pathlib import Path
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

ASSETS = Path(__file__).resolve().parent / 'assets' / 'v2'


def add_figure(doc, filename, caption, add_caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = 0
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(7.0))
    return add_caption(doc, caption)
