"""Document the October 2026 GPT-2 weight encryption experiment."""
from pathlib import Path
import json
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                x = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    x.set(qn('w:' + k), v)
                borders.append(x)
            pr.append(borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E7E6E6' if ri == 0 else ('F5F5F5' if ri % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                x = OxmlElement('w:' + edge)
                x.set(qn('w:w'), '75')
                x.set(qn('w:type'), 'dxa')
                margins.append(x)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = ri == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    return p


def add_experiment(doc, heading, body, caption):
    report_path = ROOT / 'build/gpt2_weights_demo/report.json'
    if not report_path.exists():
        report_path = ROOT / 'paper/evidence/gpt2_weights_demo/report.json'
    report = json.loads(report_path.read_text())
    assert report['result'] == 'PASS' and report['plaintext_bytes'] == 1024
    assert report['measurements']['encrypt']['cycles'] == 801
    doc.add_page_break()
    heading(doc, 'VII. WSL Ubuntu environment')
    body(doc, 'An application experiment completed on October 6, 2026 used trained GPT-2 weight bytes as the input to the existing Ascon-AEAD128 VHDL core. The workflow placed GPT-2 training in Ubuntu under Windows Subsystem for Linux 2 (WSL 2) and retained Vivado 2023.2 on Windows for RTL simulation. This allowed the Linux training process to use the laptop NVIDIA GPU without a separate desktop virtual machine or a second Vivado installation [7], [8].')
    table(doc, ['Component', 'Recorded configuration'], [
        ('Host', 'Windows 11 Home; GIGABYTE AORUS 15P YD'),
        ('CPU and memory', 'Intel Core i7-11800H; 8 cores, 16 logical processors; 32 GB RAM'),
        ('GPU', 'NVIDIA GeForce RTX 3080 Laptop GPU; 8,192 MiB VRAM'),
        ('Linux environment', 'Ubuntu 22.04.5 LTS under WSL 2; user kroix'),
        ('Python and PyTorch', 'Python 3.10; PyTorch 2.11.0+cu128'),
        ('Other libraries', 'Transformers 5.18.0; NumPy 1.26.4'),
        ('GPU driver', 'Windows NVIDIA driver 576.52; nvidia-smi reports CUDA capability 12.9'),
        ('Hardware tool', 'Vivado 2023.2 and XSim on Windows'),
    ], [1.65, 5.49])
    caption(doc, 'Table II. Software and hardware environment recorded for the GPT-2 experiment.')
    body(doc, 'Ubuntu was installed from an administrator PowerShell session with wsl --install -d Ubuntu-22.04. The distribution was confirmed as WSL version 2, and nvidia-smi inside Ubuntu detected the RTX 3080. WSL uses the Windows NVIDIA driver; a separate Linux NVIDIA display driver was not installed [8]. The cu128 suffix identifies the PyTorch CUDA 12.8 build. The CUDA version printed by nvidia-smi describes driver capability and does not by itself establish an installed CUDA compiler.')
    body(doc, 'The project was cloned into ~/projects/SAGE_Accelerator/llmc-gpt2, and a Python virtual environment named .venv kept its dependencies separate. Git, build-essential, python3-pip, and python3-venv supplied the basic development tools. GPU access was verified by allocating a 256 by 256 tensor on CUDA, multiplying it, and synchronizing the device. This calculation passed before training began. The native C/CUDA trainer was not used for GPU training; the Makefile reported that nvcc, MPI, and NCCL development support were unavailable. These components were unnecessary for the PyTorch run and CPU correctness test.')
    body(doc, 'Appendix A records the shell commands and final export code. The setup and measurements below describe this specific small experiment; they do not establish production model quality or physical FPGA performance.')

    doc.add_page_break()
    heading(doc, 'VIII. GPT-2 training and weight export')
    body(doc, 'The software was taken from SPIRE-GMU/SAGE_Accelerator, llmc-gpt2, at commit 4e586d80b61ad7577c64391202a8c43104728f1d [9]. This folder contains an llm.c implementation and the PyTorch reference trainer train_gpt2.py. The pretrained gpt2 model was loaded through Hugging Face. The run fine-tuned existing GPT-2 weights for 40 optimizer updates; it did not train GPT-2 from random initialization.')
    body(doc, 'The Tiny Shakespeare preparation script downloaded approximately 1.06 MiB of text and generated two token files. The training split contained 305,260 tokens in 611,544 bytes, and the validation split contained 32,768 tokens in 66,560 bytes. The binary sizes include a 1,024-byte dataset header. Training used the training file and validation used the separate validation file.')
    body(doc, 'Before fine-tuning, a zero-update run with batch size 4 and sequence length 64 generated the pretrained FP32 and BF16 checkpoints, tokenizer, and PyTorch debug state. The CPU test was built with make test_gpt2 and executed with OMP_NUM_THREADS=8. It passed the logits, initial loss, gradient comparisons, and all ten expected training-loss checks, ending with overall okay: 1. An earlier batch-size-1 debug file produced correct tensor comparisons but failed the fixed training-loss expectations; regenerating the reference with the standard batch size 4 resolved the mismatch. The zero-update timing warning nanms only reflected an empty timing list.')
    table(doc, ['Training setting', 'Value'], [
        ('Compute and precision', 'One RTX 3080; FP32; compilation disabled'),
        ('Microbatch and sequence', '1 sequence of 64 tokens'),
        ('Effective batch', '256 tokens; 4 gradient accumulation steps'),
        ('Optimizer updates', '40; overfit_single_batch disabled'),
        ('Learning rate', '0.0001; no warmup; default decay fraction 1.0'),
        ('Validation sampling', 'Every 10 updates and at completion; 5 batches per check'),
        ('Output directory', 'runs/ascon_demo_40'),
    ], [2.0, 5.14])
    caption(doc, 'Table III. Settings for the GPU fine-tuning demonstration.')
    body(doc, 'Validation loss was 5.404110 before training and 3.810759 after 40 updates. Intermediate checks after 10, 20, and 30 updates were 4.078157, 3.942588, and 3.860151. Each value averages five validation batches rather than the entire validation split. The trainer reported a mean of 132.058 ms over its final 19 recorded iteration timings and peak PyTorch allocated GPU memory of 2,416 MiB. These iteration timings are not the total elapsed run time and include any validation work scheduled in a timed iteration.')
    body(doc, 'The supplied PyTorch bridge exports initialization weights before its training loop. An additional export block was therefore added after the loop to call write_model(raw_model, ..., dtype="float32") and save the updated weights as gpt2_trained_fp32.bin. The file size was 497,904,640 bytes: a 1,024-byte header plus 497,903,616 bytes of FP32 parameter data, including vocabulary padding from 50,257 to 50,304 entries. This is approximately 474.84 MiB. The serialized payload contains 124,475,904 FP32 slots; the unpadded model has 124,439,808 unique parameters.')
    body(doc, 'A 1 KiB sample was extracted by skipping the checkpoint header and reading the next 1,024 bytes, equivalent to 256 FP32 values from the first token embedding. The file was verified against that trained payload, and the trained checkpoint was confirmed to differ from the pretrained checkpoint. The sample was copied to Windows as gpt2_trained_weights_1KiB.bin. It represents only a small portion of model weights and cannot reconstruct the model on its own.')

    doc.add_page_break()
    heading(doc, 'IX. Ascon encryption of trained weight bytes')
    body(doc, 'The original official-vector TextIO testbench stores only two 16-byte payload blocks, so it cannot accept the entire 1 KiB sample. A separate runner, verification/gpt2_weights_demo.py, generated a dedicated tb_gpt2_weights testbench and an isolated Vivado project under build/gpt2_weights_demo. The runner reused the existing ascon_pkg.vhd, ascon_permutation.vhd, and ascon_aead128_core.vhd. The application test streamed all 64 full 16-byte beats as one AEAD message; the original RTL and official-vector testbench were not altered.')
    body(doc, 'The Python Ascon reference first passed all 1,089 pinned official known-answer vectors. It then computed the expected ciphertext and tag for the actual trained-weight sample. The testbench checked every ciphertext block and the tag, then decrypted the ciphertext and checked all recovered bytes and successful authentication. Full block dumps were preserved in full_output.txt, while simulation_results.txt and simulation_vectors.txt recorded the PASS results and cycle counts. The runner verified the complete byte arrays before producing report.json.')
    table(doc, ['Measurement or input', 'Recorded result'], [
        ('Plaintext and associated data', '1,024 weight bytes; 0 associated-data bytes'),
        ('Demo key', '000102030405060708090a0b0c0d0e0f'),
        ('Demo nonce', '101112131415161718191a1b1c1d1e1f'),
        ('Clock', '100 MHz; 10 ns per cycle'),
        ('Encryption', '801 command-to-done cycles; 8.01 microseconds; PASS'),
        ('Decryption', '801 command-to-done cycles; 8.01 microseconds; PASS'),
        ('Payload throughput', '1,022.72 Mb/s calculated from simulated latency'),
        ('Authentication tag', 'e3a4ec0147a1f6681ba8f8aeb188a49f'),
        ('Ciphertext and tag sizes', '1,024-byte ciphertext and a separate 16-byte tag'),
        ('Round-trip integrity', 'Recovered 1,024 bytes exactly match the input sample'),
    ], [2.0, 5.14])
    caption(doc, 'Table IV. Behavioral XSim results for one 1 KiB trained-weight sample.')
    body(doc, 'The cycle counter was sampled when a ready command was offered and again when done was observed. Thus the interval includes command submission, initialization, the empty associated-data phase, all payload handshakes, finalization, tag handling, and completion. The driver accepted each output beat before submitting the next input beat. The measurement uses that testbench schedule and must not be described as an ideal unstalled datapath limit. The latency calculation is 801 cycles times 10 ns per cycle, giving 8,010 ns or 8.01 microseconds. Payload throughput is 1,024 times 8 bits divided by 8.01 microseconds.')
    body(doc, 'The source sample SHA-256 is recorded below. The Windows copy matched the Ubuntu export, binding the encryption result to the actual trained-weight bytes. The demo key and nonce are public test inputs. A deployed encryption system would require protected keys and unique nonces for each message under a given key; repeating this fixed pair is appropriate only for reproducing this test.')
    code(doc, 'fa00e0b873c026e9c25298fc1428872a0968\n0e11e0ec5ac74af29fcab3e34541')
    body(doc, 'The 8.01-microsecond result is functional RTL simulation at a chosen 100 MHz clock. No physical board encryption, UART/USB transfer, DMA, PC file I/O, or post-route application timing was measured. The small-payload result of 352.441 Mb/s in Section VI averages different associated-data lengths, whereas this case has empty associated data and a much larger payload. The numbers therefore describe different workloads and should not be treated as a controlled architectural speedup.')


def add_appendix(doc, heading, body):
    doc.add_page_break()
    heading(doc, 'Appendix A Reproducing the environment and training')
    body(doc, 'Commands are grouped by the shell in which they run. WSL installation uses administrator PowerShell; the project and training commands run in Ubuntu. The following commands document the completed setup. Package versions observed in Table II should be retained when reproducing the numerical results, because an unpinned requirements installation may select newer releases.', indent=False)
    heading(doc, 'A1 Windows PowerShell setup')
    code(doc, 'wsl --install -d Ubuntu-22.04\n# Restart Windows if requested, then continue:\nwsl --update\nwsl --list --verbose\nwsl -d Ubuntu-22.04')
    body(doc, 'Create the Ubuntu username and password on first launch. Confirm VERSION 2 in the distribution list. Password characters are intentionally invisible at the prompt. WSL GPU access was then checked inside Ubuntu with nvidia-smi.', indent=False)
    heading(doc, 'A2 Ubuntu project and Python environment')
    code(doc, 'cd ~\nsudo apt update\nsudo apt install -y git build-essential python3-pip python3-venv\nmkdir -p ~/projects\ncd ~/projects\ngit clone https://github.com/SPIRE-GMU/SAGE_Accelerator.git\ncd SAGE_Accelerator\ngit checkout 4e586d80b61ad7577c64391202a8c43104728f1d\ncd llmc-gpt2\npython3 -m venv .venv\nsource .venv/bin/activate\npython -m pip install --upgrade pip\npython -m pip install torch --index-url https://download.pytorch.org/whl/cu128\npython -m pip install -r requirements.txt\nnvidia-smi')
    body(doc, 'The checkout command above pins the source revision used in this experiment. For later sessions, return to llmc-gpt2 and reactivate .venv. No Linux NVIDIA display driver or separate CUDA compiler installation was needed for the PyTorch workflow [8], [10].', indent=False)
    heading(doc, 'A3 CUDA calculation check')
    code(doc, 'python3 - <<\'PY\'\nimport torch\nprint(torch.__version__)\nassert torch.cuda.is_available()\nprint(torch.cuda.get_device_name(0))\nx = torch.randn(256, 256, device="cuda")\ny = x @ x\ntorch.cuda.synchronize()\nprint(y.shape)\nPY')
    body(doc, 'The closing PY must be on a line by itself. During setup, visible bracketed-paste markers were resolved by typing bind \'set enable-bracketed-paste off\' in Bash. This changes terminal input behavior and has no effect on model results.', indent=False)

    doc.add_page_break()
    heading(doc, 'Appendix A Dataset and correctness checks')
    code(doc, 'python3 dev/data/tinyshakespeare.py\npython3 train_gpt2.py --device cuda --model gpt2 \\\n  --num_iterations 0 --write_tensors 1 --batch_size 4 \\\n  --sequence_length 64 --total_batch_size 256\nmake test_gpt2\nOMP_NUM_THREADS=8 ./test_gpt2')
    body(doc, 'Require overall okay: 1 from the C test. Use batch size 4 for this reference because the ten expected losses in test_gpt2.c are fixed to that workload. A zero-update reference run reports nanms because it has no training durations to average.', indent=False)
    heading(doc, 'A4 Final weight export added after the training loop')
    body(doc, 'Append the following block at the end of train_gpt2.py, inside its main block. Its first line has four leading spaces. It saves updated FP32 weights and selects actual weight bytes after the 1,024-byte model header.', indent=False)
    code(doc, '''    if master_process and args.num_iterations > 0:
        export_dir = args.output_dir or "."
        os.makedirs(export_dir, exist_ok=True)
        export_path = os.path.join(export_dir, "gpt2_trained_fp32.bin")
        write_model(raw_model, export_path, dtype="float32")
        file_bytes = os.path.getsize(export_path)
        print0(f"Trained checkpoint: {file_bytes} bytes")
        print0(f"Weight payload including vocabulary padding: {file_bytes - 1024} bytes")
        with open(export_path, "rb") as source:
            source.seek(1024)
            sample = source.read(1024)
        assert len(sample) == 1024
        with open(os.path.join(export_dir, "weights_1KiB.bin"), "wb") as target:
            target.write(sample)
        print0("Exported 1024 weight bytes for Ascon")''')
    code(doc, 'python3 -m py_compile train_gpt2.py')
    body(doc, 'The export is guarded by a positive iteration count so a reference-only run cannot be mislabeled as trained weights. The original before-training bridge files remain distinct from the final trained checkpoint.', indent=False)

    doc.add_page_break()
    heading(doc, 'Appendix A Training transfer and XSim execution')
    heading(doc, 'A5 Ubuntu GPU fine-tuning command')
    code(doc, 'python3 train_gpt2.py \\\n  --device cuda --model gpt2 \\\n  --input_bin dev/data/tinyshakespeare/tiny_shakespeare_train.bin \\\n  --input_val_bin dev/data/tinyshakespeare/tiny_shakespeare_val.bin \\\n  --batch_size 1 --sequence_length 64 --total_batch_size 256 \\\n  --num_iterations 40 --overfit_single_batch 0 \\\n  --val_loss_every 10 --val_max_steps 5 --write_tensors 0 \\\n  --output_dir runs/ascon_demo_40')
    heading(doc, 'A6 Inspect and transfer the exported sample')
    code(doc, 'stat -c "%n %s bytes" runs/ascon_demo_40/gpt2_trained_fp32.bin\nstat -c "%n %s bytes" runs/ascon_demo_40/weights_1KiB.bin\nsha256sum runs/ascon_demo_40/weights_1KiB.bin\ncp -i runs/ascon_demo_40/weights_1KiB.bin \\\n  /mnt/c/Users/Symok/Downloads/gpt2_trained_weights_1KiB.bin')
    heading(doc, 'A7 Windows PowerShell VHDL simulation')
    code(doc, 'cd "C:\\Users\\Symok\\Downloads\\ascon-aead128-vhdl-main"\npython verification/gpt2_weights_demo.py `\n  "C:\\Users\\Symok\\Downloads\\gpt2_trained_weights_1KiB.bin"')
    body(doc, 'The runner uses Vivado on PATH or the installed C:\\Xilinx\\Vivado\\2023.2\\bin\\vivado.bat. It first validates the Python oracle and then invokes XSim using a generated testbench. The project and results are isolated under build/gpt2_weights_demo. Before rerunning, archive that output directory if the previous receipts must be retained.', indent=False)
    table(doc, ['Artifact', 'Purpose'], [
        ('report.json and inputs.json', 'Inputs, sample hash, timing definition, and verified result'),
        ('simulation_results.txt', 'VHDL ENCRYPT and DECRYPT PASS lines and 801-cycle counts'),
        ('full_output.txt', 'All 64 output blocks in each direction'),
        ('ciphertext.bin and tag.bin', '1,024-byte ciphertext and 16-byte authentication tag'),
        ('recovered_weights.bin', 'Byte-exact decryption result'),
        ('vivado_console.log', 'Compile, elaboration, and simulation receipts'),
        ('runs/ascon_demo_40/main.log', 'Ubuntu training and validation losses'),
    ], [2.5, 4.64])
    body(doc, 'The training output, source revision, installed software versions, sample hash, and simulation receipts jointly establish the provenance of this experiment. Reported training memory is PyTorch allocation; reported encryption time is simulation latency. Neither is a physical board measurement.', indent=False)
def add_latency_research(doc, heading, body, caption):
    doc.add_page_break()
    heading(doc, 'X. Reducing latency at a fixed clock frequency')
    body(doc, 'The 1 KiB experiment provides a measured simulation baseline for evaluating architectural changes: 801 cycles, or 8.01 microseconds at 100 MHz. The following options are proposed research directions; none has been implemented or measured in this experiment. Published Ascon hardware supports multiple rounds per cycle, including two- and four-round configurations [11]. OpenTitan also identifies multiple rounds per cycle as a way to reduce permutation latency [12]. These sources establish the technique, not a performance result for this FPGA design.')
    heading(doc, 'A. Round unrolling and conditional cycle estimates')
    body(doc, 'The current permutation engine performs one round per cycle. With no associated data, the 64 full message blocks require 64 eight-round permutations, in addition to twelve-round initialization and finalization. This accounts for 12 + 64 × 8 + 12 = 536 round-computation cycles. The remaining 265 cycles include controller transitions, input/output handshakes, and completion work. They cannot all be assumed removable.')
    body(doc, 'Round unrolling places consecutive rounds in the combinational path between state registers. Two rounds per cycle would reduce the round-computation portion to 268 cycles; four would reduce it to 134. Holding the remaining 265 cycles unchanged gives the estimates below. Each estimate assumes that the larger combinational circuit still meets the 10 ns clock period and that controller behavior is otherwise unchanged.')
    table(doc, ['Architecture at 100 MHz', 'Total cycles', 'Latency', 'Reduction'], [
        ('Current one-round core: measured in XSim', '801', '8.01 µs', 'Baseline'),
        ('Two rounds per cycle: conditional estimate', '533', '5.33 µs', '33.5%'),
        ('Four rounds per cycle: conditional estimate', '399', '3.99 µs', '50.2%'),
    ], [3.6, 1.05, 1.15, 1.34])
    caption(doc, 'Conditional architecture estimates for the same 1 KiB message and empty associated data. Only the baseline is a simulation result; optimized timing and resource use remain unverified.')
    body(doc, 'Two rounds per cycle is the first candidate to evaluate because it offers a substantial cycle reduction with a smaller increase in logic depth than four rounds. Unrolling increases logic and routing demands, so a lower cycle count is useful only if routed timing still meets 100 MHz. The estimates are derived from this workload and must not be presented as published benchmark measurements.')
    heading(doc, 'B. Other ways to reduce waiting')
    table(doc, ['Option', 'Expected benefit and practical limit'], [
        ('Overlap output transfer and permutation', 'A registered output buffer can retain the completed ciphertext while the next permutation runs. The current controller waits for output acceptance before starting that permutation. Benefit depends on stalls; ready/valid behavior must remain correct.'),
        ('Reduce controller handoff cycles', 'Inspect the start, wait, and completion transitions for avoidable idle cycles. A cycle trace is needed before assigning a numerical saving.'),
        ('Multiple cores for independent chunks', 'Separate authenticated chunks can run concurrently, improving completion time for large files. Each chunk needs a unique nonce and tag, with chunk order and model identity authenticated. Transfer bandwidth can limit the benefit.'),
        ('Pipeline registers within a round', 'Can shorten the combinational path, but at a fixed clock may add single-message cycles. Consecutive rounds and message blocks depend on the previous state; independent messages are needed to exploit interleaving.'),
    ], [2.05, 5.09])
    body(doc, 'The existing 128-bit input interface already matches the Ascon-AEAD128 message rate. Widening the external bus alone does not remove the sequential permutation dependency. The standard round counts must remain unchanged; the optimization is to perform the required rounds more efficiently, not to reduce cryptographic work.')
    heading(doc, 'C. Proposed comparison with the baseline')
    body(doc, 'The next study should record a baseline cycle trace, implement a two-round permutation variant, and repeat the official-vector tests and the same 1 KiB encryption/decryption experiment. Report cycle latency, routed timing at 100 MHz, LUT and register use, and estimated power for both designs. Simulation, Vivado implementation estimates, and physical-board measurements should be reported separately. This comparison would determine whether the predicted cycle reduction is achieved without compromising functional correctness or timing closure.')
