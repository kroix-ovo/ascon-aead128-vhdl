#!/usr/bin/env python3
"""Build the editable undergraduate research paper as a DOCX file."""

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from gpt2_report_sections import add_experiment, add_appendix, add_latency_research
from behavior_report_sections import add_behavior_results
from weight_stream_report_sections import add_weight_stream_results
from fold_report_sections import add_fold_results


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "paper" / "VHDL_Implementation_of_NIST_Ascon_AEAD128_on_a_Nexys_A7_FPGA.docx"
METRICS = ROOT / "paper" / "evidence" / "baseline_20260918" / "metrics.json"


def load_metrics():
    source = METRICS if METRICS.exists() else ROOT / "paper/evidence/baseline_20260918/metrics.json"
    if not source.exists():
        return {"simulation": {"complete": False}, "vivado": {"complete": False}}
    return json.loads(source.read_text(encoding="utf-8"))


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, values in edges.items():
        tag = "w:" + edge_name
        edge = borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            borders.append(edge)
        for key, value in values.items():
            edge.set(qn("w:" + key), str(value))


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_body(doc, text, indent=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.02
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.22)
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = False
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    return paragraph


def add_result_table(doc, metrics):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.05), Inches(1.45), Inches(1.25), Inches(2.05)]
    headings = ["Verification level", "Cases", "Result", "Evidence"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(headings[index])
        r.bold = True
        r.font.size = Pt(8.2)
    textio_complete = metrics.get("simulation", {}).get("complete", False)
    rows = [
        ("Python reference", "81 boundary pairs", "PASS", "official KAT comparison"),
        ("GHDL direct VHDL", "1,089 × 2", "PASS", "encrypt and decrypt"),
        ("Protocol/security", "9 test groups", "PASS", "stalls, reset, zeroization"),
        (
            "GHDL TextIO",
            "1,089 × 2" if textio_complete else "not rerun",
            "PASS" if textio_complete else "PENDING",
            "two output text files",
        ),
        (
            "Vivado XSim",
            "1,089 × 2" if textio_complete else "not run",
            "PASS" if textio_complete else "PENDING",
            "official KAT TextIO bench",
        ),
        ("Icarus netlist", "1,089 × 2", "PASS", "GHDL-generated Verilog"),
        ("Verilator netlist", "1,089 × 2", "PASS", "GHDL-generated Verilog"),
        (
            "Vivado implementation",
            "placed and routed" if metrics.get("vivado", {}).get("complete") else "not run",
            "PASS" if metrics.get("vivado", {}).get("complete") else "PENDING",
            "reports and bitstream" if metrics.get("vivado", {}).get("complete") else "requires Windows/Linux host",
        ),
        (
            "Nexys A7 board",
            "four demo vectors",
            "PENDING",
            "programming and hardware test",
        ),
    ]
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.size = Pt(8.0)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "single", "sz": 5, "color": "000000"},
                bottom={"val": "single", "sz": 5, "color": "000000"},
                left={"val": "single", "sz": 5, "color": "000000"},
                right={"val": "single", "sz": 5, "color": "000000"},
            )
    return table


def build():
    metrics = load_metrics()
    vivado = metrics.get("vivado", {})
    if vivado.get("complete"):
        implementation_summary = (
            f"Vivado used {int(vivado['slice_luts'])} slice LUTs and "
            f"{int(vivado['slice_registers'])} slice registers, with no block RAM. "
            f"The routed design met the 100 MHz constraint with "
            f"{vivado['wns_ns']:.3f} ns worst setup slack. Vivado estimated "
            f"{vivado['total_on_chip_power_w']:.3f} W total on-chip power and "
            "generated the FPGA bitstream. Physical board programming and "
            "demonstration testing were pending at this baseline stage; Section XIII records the later board tests."
        )
    else:
        implementation_summary = (
            "Vivado implementation measurements remain pending because Vivado is not "
            "available on the Mac used for this stage of the work."
        )
    simulation = metrics.get("simulation", {})
    latency = metrics.get("latency_by_payload_bytes", {}).get("encryption", {})
    if simulation.get("complete") and "32" in latency:
        latency_32 = latency["32"]
        latency_summary = (
            f"The official-vector TextIO run measured 34 to 88 command-to-done "
            f"cycles. For a 32-byte payload, the mean across all 33 associated-data "
            f"lengths was {latency_32['mean_cycles']:.3f} cycles, equal to "
            f"{latency_32['mean_payload_throughput_mbps']:.3f} Mb/s at 100 MHz. "
            "These are cycle-based simulation results, not a measured maximum clock rate."
        )
    else:
        latency_summary = "Cycle-count results are generated after the full TextIO run."

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Jones | ")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title_borders = OxmlElement("w:pBdr")
    title_bottom = OxmlElement("w:bottom")
    title_bottom.set(qn("w:val"), "nil")
    title_borders.append(title_bottom)
    title._p.get_or_add_pPr().append(title_borders)
    title_run = title.add_run("VHDL Implementation of NIST Ascon AEAD128\non a Nexys A7 FPGA")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(7)
    author.add_run(
        "Kroix Jones\nElectrical Engineering Undergraduate Researcher\n"
        "Howard University\nResearch Advisor: Professor Wenjie Che"
    )

    add_heading(doc, "Abstract")
    add_body(
        doc,
        "This paper describes a VHDL-2008 implementation of the standardized NIST "
        "Ascon-AEAD128 algorithm for a Nexys A7 100T FPGA. The design uses one shared "
        "320-bit permutation datapath with configurable round unrolling and optional input/output queues. A "
        "128-bit streaming interface accepts associated data and payload bytes, including "
        "empty and partial blocks. Encryption returns ciphertext and a tag. Decryption "
        "returns provisional plaintext that an external quarantine releases only after tag verification. The "
        "portable test flow passed all 1,089 official known-answer vectors in both "
        "directions with GHDL and Vivado XSim. Generated Verilog was also checked with "
        "Icarus Verilog and Verilator. Vivado 2023.2 completed synthesis, placement, "
        "routing, timing analysis, and bitstream generation. At a simulated 100 MHz, "
        "the complete 2,397,696-byte FP32 adapter payload encrypts and decrypts in "
        "449,581 cycles per direction with folded absorption, compared with "
        "899,153 for the buffered three-round core and 1,798,305 for the one-round "
        "baseline. All recovered bytes match and authentication passes. The folded "
        "core routes at 100 MHz with unconstrained external I/O; a board-integrated "
        "copy also meets 100 MHz in one run. Earlier physical captures validate only "
        "the buffered design. The folded core has not been programmed or measured "
        "on hardware. Section XIV separates these evidence scopes and documents "
        "the later AD padding and decryption tag output fixes.",
        indent=False,
    )

    add_heading(doc, "I. Research objective")
    add_body(
        doc,
        "The goal of this work is to build an understandable FPGA implementation of "
        "authenticated encryption. The core must follow the final standard, handle normal "
        "streaming stalls, and make authentication failure visible to the surrounding "
        "system. The first version favors a clear control path over maximum throughput. "
        "It is meant to support later research on faster datapaths and hardware security "
        "countermeasures without changing the external transaction format.",
    )

    add_heading(doc, "II. Ascon-AEAD128 background")
    add_body(
        doc,
        "NIST SP 800-232 defines Ascon-AEAD128 with a 128-bit key, 128-bit nonce, "
        "128-bit authentication tag, 128-bit data rate, and a 320-bit internal state [1]. "
        "The state is divided into five 64-bit words. Initialization and finalization use "
        "a 12-round permutation. Associated-data and full payload blocks use an 8-round "
        "permutation. Each round applies a constant to one state word, evaluates 64 "
        "five-input S-box slices in parallel, and then performs word-specific rotations "
        "and XOR operations.",
    )
    add_body(
        doc,
        "The final NIST version was selected instead of legacy Ascon-128 or Ascon-128a. "
        "The final standard changed the initialization value and changed byte mapping to "
        "little endian [1]. For this core, byte lane 0 is bits 7 through 0 of the 128-bit "
        "port. Padding places the byte value 01 immediately after the last valid byte. "
        "Using the official mapping avoids a common result where a design is internally "
        "consistent but does not match current NIST vectors.",
    )

    flow = doc.add_table(rows=1, cols=6)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["Initialize\np12", "Absorb AD\np8", "Domain\nseparate", "Process data\np8", "Finalize\np12", "Tag\n128 bits"]
    for idx, cell in enumerate(flow.rows[0].cells):
        cell.width = Inches(1.18)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(labels[idx])
        r.font.size = Pt(8.3)
        if idx in (0, 5):
            r.bold = True
        set_cell_border(
            cell,
            top={"val": "single", "sz": 8, "color": "000000"},
            bottom={"val": "single", "sz": 8, "color": "000000"},
            left={"val": "single", "sz": 8, "color": "000000"},
            right={"val": "single", "sz": 8, "color": "000000"},
        )
    add_caption(doc, "Fig. 1. Ascon-AEAD128 processing phases, adapted from NIST SP 800-232 [1].")

    doc.add_page_break()
    add_heading(doc, "III. Original hardware architecture")
    add_body(
        doc,
        "The datapath contains five 64-bit state registers and one combinational round. "
        "A small round counter selects the required constants and returns the updated state "
        "after 8 or 12 clocks. Reusing one round reduces duplicated logic and keeps the "
        "relationship between the VHDL and the algorithm visible. The same engine serves "
        "initialization, associated data, payload processing, and finalization.",
    )
    pic = doc.add_picture(str(ROOT / "paper" / "assets" / "ascon_architecture.png"), width=Inches(7.0))
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Fig. 2. Original block diagram of the implemented architecture.")
    add_body(
        doc,
        "The controller accepts a command containing the key, nonce, and direction. "
        "Associated data must finish before payload begins. Each stream beat includes 16 "
        "keep bits and a last flag. Keep bits must start at lane 0 and remain contiguous. "
        "An empty phase uses a terminal beat with no keep bits set. Any invalid ordering or "
        "mask produces a protocol error and sends all transaction storage through the "
        "same zeroization state used after a normal result.",
    )

    add_heading(doc, "IV. Authentication-safe decryption")
    add_body(
        doc,
        "A decryption core cannot know whether plaintext is authentic until it processes "
        "the final tag. Storing every possible message inside the cryptographic core would "
        "make its memory size depend on the application. This design sends tentative "
        "plaintext to an external quarantine buffer instead. A correct tag pulses commit. "
        "A wrong tag pulses discard and never pulses commit. The interface therefore makes "
        "the safe release rule explicit, but the system integrator must enforce it.",
    )

    doc.add_page_break()
    add_heading(doc, "V. Original controller and verification")
    add_body(
        doc,
        "The finite-state machine was made for this streaming architecture. Start states "
        "issue a one-cycle request to the permutation engine. Wait states prevent another "
        "request until the current permutation ends. Output states hold data stable during "
        "backpressure. Tag states either wait for the encryption tag to be accepted or wait "
        "for a decryption tag to be supplied.",
    )
    pic = doc.add_picture(str(ROOT / "paper" / "assets" / "ascon_fsm.png"), width=Inches(7.0))
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Fig. 3. Original controller state flow used by the VHDL core.")
    add_body(
        doc,
        "Verification began with an independent Python model. The model was compared with "
        "the official ascon-c known-answer file pinned to a recorded source revision [2]. "
        "Cocotb then drove the same vectors through the VHDL core. The full run covered all "
        "1,089 combinations of associated-data and plaintext lengths from 0 through 32 "
        "bytes, with encryption followed by decryption. Separate tests stalled each output, "
        "reset an active transaction, sent invalid masks and phase ordering, checked fixed-"
        "seed messages up to 80 bytes, and changed the key, nonce, associated data, ciphertext, "
        "and tag.",
    )
    add_result_table(doc, metrics)
    add_caption(doc, "Table I. Verification and implementation status as of September 18, 2026.")

    doc.add_page_break()
    add_heading(doc, "VI. Vivado flow and FPGA demonstration")
    add_body(
        doc,
        "The repository contains Tcl scripts for Vivado 2023.2 and an XDC file for the "
        "Nexys A7 100T part xc7a100tcsg324-1. The constraint sets the board clock to 100 "
        "MHz. A pure VHDL XSim testbench writes simulation_results.txt for human review "
        "and simulation_vectors.txt for scripts. The same file-driven bench was checked "
        "with GHDL across all 1,089 vectors in both directions. XSim used the identical "
        "TextIO input and completed 2,178 encryption and decryption checks with no "
        "failures. " + latency_summary + " " + (
            "The implementation results were extracted from the routed Vivado reports. "
            + implementation_summary + " The original routed report included a CFGBVS-1 "
            "configuration warning. The repository XDC now sets CFGBVS to VCCO and "
            "CONFIG_VOLTAGE to 3.3 for the Nexys A7 configuration bank. Physical-board "
            "validation was pending at this baseline stage; later results appear in Section XIII. The power value is a Vivado estimate rather "
            "than a board measurement."
            if vivado.get("complete")
            else "XSim execution, synthesis, implementation, timing, and power reporting "
            "still require the Vivado machine. Resource use, maximum clock rate, and "
            "post-route throughput are intentionally not estimated in this paper."
        ),
    )
    add_body(
        doc,
        "The board wrapper stores four official vectors: empty input, a one-byte partial "
        "block, one full block, and a 17-byte two-block case. Switches choose the vector, "
        "encryption or decryption, payload or tag display, block number, and byte number. "
        "The center button starts a run and the upper button resets the design. LEDs display "
        "the selected result byte, busy and done state, authentication pass or fail, and a "
        "protocol error.",
    )

    add_experiment(doc, add_heading, add_body, add_caption)
    add_latency_research(doc, add_heading, add_body, add_caption)
    add_behavior_results(doc, add_heading, add_body, add_caption)
    add_weight_stream_results(doc, add_heading, add_body, add_caption)
    add_fold_results(doc, add_heading, add_body, add_caption)
    doc.add_page_break()
    add_heading(doc, "XV. Limitations and next work")
    add_body(
        doc,
        "This version is an unprotected cryptographic core. It does not prevent power or "
        "electromagnetic side-channel analysis, and it does not detect injected faults. "
        "Kandi et al. describe threshold and redundancy-based hardware options that could "
        "be studied in subsequent work [3]. Sections XIII and XIV measure the implemented unrolling "
        "and buffering tradeoff, including folded absorption. Independent lanes, memory transport and lossless compression "
        "remain separate experiments requiring additional interfaces and components. The external quarantine buffer is also a "
        "required part of safe decryption, not an optional application feature.",
    )

    add_heading(doc, "XVI. Findings")
    add_body(
        doc,
        "For this adapter, eliminating controller handoffs halves the buffered "
        "three-round core's cycle count without removing rounds. The folded core "
        "requires 449,581 cycles per direction, matches the reference ciphertext and "
        "tag, and reconstructs the complete checkpoint exactly. Its three-clock "
        "full-block interval is supported by simulation; its 100 MHz closure is a "
        "routed implementation result. Physical full-file transfer, memory bandwidth "
        "and authenticated compute integration remain unmeasured. Both the AD "
        "empty-terminator defect and legacy decryption tag exposure are fixed in "
        "this branch, with focused regressions and negative controls.",
    )

    add_heading(doc, "References")
    references = [
        "[1] M. S. Turan, K. A. McKay, D. Chang, J. Kang, and J. Kelsey, “Ascon-Based Lightweight Cryptography Standards for Constrained Devices,” NIST SP 800-232, Aug. 2025, doi: 10.6028/NIST.SP.800-232.",
        "[2] Ascon Team, “ascon-c: Reference and optimized implementations of Ascon,” commit 446347f21b209f3921c65ece70027c366cbe1693. [Online]. Available: https://github.com/ascon/ascon-c",
        "[3] A. Kandi et al., “Hardware Implementation of ASCON,” NIST Lightweight Cryptography Workshop, June 2023. https://csrc.nist.gov/csrc/media/Events/2023/lightweight-cryptography-workshop-2023/documents/accepted-papers/07-hardware-implementation-of-ascon.pdf",
        "[4] K. Liu, L. Han, and Z. Quan, “A High Performance Ascon Architecture for Data Encryption and Authentication,” in Proc. 5th Int. Conf. on Cryptography, Network Security and Communication Technology (CNSCT 2026), Haikou, China, Jan. 23–25, 2026, doi: 10.1145/3802927.3802941.",
        "[5] N. T. M. Huong, T. V. Nam, and P. V. Phat, “Application of the Ascon-AEAD128 Algorithm for IoT Data Security,” 11th Int. Conf. on Applying New Technology in Green Buildings, 2026, doi: 10.1109/ATIGB70203.2026.11628316.",
        "[6] Digilent, “Nexys A7-100T Master XDC.” [Online]. Available: https://github.com/Digilent/digilent-xdc",
        "[7] Microsoft, “Install WSL.” [Online]. Available: https://learn.microsoft.com/en-us/windows/wsl/install. Accessed Oct. 6, 2026.",
        "[8] NVIDIA, “CUDA on WSL User Guide.” [Online]. Available: https://docs.nvidia.com/cuda/wsl-user-guide/. Accessed Oct. 6, 2026.",
        "[9] SPIRE-GMU, “SAGE_Accelerator, llmc-gpt2,” commit 4e586d80b61ad7577c64391202a8c43104728f1d. [Online]. Available: https://github.com/SPIRE-GMU/SAGE_Accelerator/tree/4e586d80b61ad7577c64391202a8c43104728f1d/llmc-gpt2.",
        "[10] PyTorch, “Start Locally.” [Online]. Available: https://pytorch.org/get-started/locally/. Accessed Oct. 6, 2026.",
        "[11] R. Primas, “Hardware Design of Ascon (SP800-232),” ascon-verilog. [Online]. Available: https://github.com/rprimas/ascon-verilog. Accessed Oct. 6, 2026.",
        "[12] OpenTitan, “Ascon: Theory of Operation.” [Online]. Available: https://opentitan.org/book/hw/ip/ascon/doc/theory_of_operation.html. Accessed Oct. 6, 2026.",
        "[13] C. E. LaForest, “Rules for Ready/Valid Handshakes and Synchronization,” FPGA Design Elements. [Online]. Available: https://fpgacpu.ca/fpga/handshake.html. Accessed Oct. 21, 2026.",
        "[14] K. Jones, Folded Ascon weight stream study, October 21, 2026. Repository reports/Ascon weight stream bottlenecks.md and docs/evidence/handoff_study_fold_20261021/. https://github.com/kroix-ovo/ascon-aead128-vhdl/tree/ascon-weight-stream-fold-absorb",
        "[15] T. Itschner, ascon_cipher_verilog, source architecture reference. https://github.com/tristanitschner/ascon_cipher_verilog. No third-party timing result is adopted here.",
        "[16] K. Jones, Associated data padding and decryption tag output fixes, October 21, 2026. Repository docs/ad_padding_fix.md; docs/evidence/ad_padding_fix_20261021/; docs/evidence/tag_output_gate_20261021/.",
        "[17] AMD, Vivado Design Suite User Guide Design Analysis and Closure Techniques, UG906, Category 1 Timing, 2026.1. https://docs.amd.com/r/en-US/ug906-vivado-design-analysis/Category-1-Timing",
        "[18] W. Hua, M. Umar, Z. Zhang and G. E. Suh, GuardNN Secure Accelerator Architecture for Privacy-Preserving Deep Learning, DAC 2022, doi: 10.1145/3489517.3530439. https://arxiv.org/html/2008.11632v2.",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(1.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(ref)
        r.font.size = Pt(8.2)

    add_appendix(doc, add_heading, add_body)
    core = doc.core_properties
    core.title = "VHDL Implementation of NIST Ascon AEAD128 on a Nexys A7 FPGA"
    core.author = "Kroix Jones"
    core.subject = "Undergraduate research with Professor Wenjie Che"
    core.keywords = "Ascon-AEAD128, VHDL, FPGA, Nexys A7, authenticated encryption"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
